import openpyxl
import json
import os
import importlib.util

from constants import PAGE_METADATA, TABS_METADATA


def voices_tab_big_numbers(excel_file):
    try:
        print("\n🔵 Starting Voices Tab Big Numbers Update")

        # --------------------------------------------------
        # Paths
        # --------------------------------------------------
        script_dir = os.path.dirname(os.path.abspath(__file__))
        json_path = os.path.join(
            script_dir, "..", "pages", "voices-from-the-ground.json"
        )

        print(f"📊 Excel File: {excel_file}")
        print(f"📄 JSON Path: {json_path}")

        # --------------------------------------------------
        # Load Excel
        # --------------------------------------------------
        workbook = openpyxl.load_workbook(excel_file, data_only=True)

        sheet_name = PAGE_METADATA["VOICE_TAB_BIG_NUMBERS"]
        if sheet_name not in workbook.sheetnames:
            print("❌ Sheet not found:", sheet_name)
            print("Available:", workbook.sheetnames)
            return

        sheet = workbook[sheet_name]
        print(f"✅ Using Excel sheet: {sheet_name}")

        # --------------------------------------------------
        # Read headers (PRESERVE INDEX)
        # --------------------------------------------------
        raw_headers = [cell.value for cell in sheet[1]]
        cleaned_headers = [
            str(h).strip() if h is not None else ""
            for h in raw_headers
        ]

        print("\n🧾 Headers with index:")
        for i, h in enumerate(cleaned_headers):
            if h:
                print(f"   [{i}] {h}")

        # --------------------------------------------------
        # Indicator headers start from index 2
        # --------------------------------------------------
        indicator_headers = [
            str(col).strip()
            for col in TABS_METADATA["VOICE_TAB_BIG_NUMBERS"][2:]
        ]

        if not all(col in cleaned_headers for col in indicator_headers):
            print("\n❌ Missing required indicator columns")
            print("Expected:", indicator_headers)
            print("Found:", cleaned_headers)
            return

        print("\n✅ All indicator columns are present")

        # --------------------------------------------------
        # Helper: Sum column as INTEGER
        # --------------------------------------------------
        def sum_by_header(header_name):
            idx = cleaned_headers.index(header_name)
            total = 0
            for row in sheet.iter_rows(min_row=2):
                val = row[idx].value
                if isinstance(val, (int, float)):
                    total += int(val)
            return total

        # --------------------------------------------------
        # Calculate indicator totals
        # --------------------------------------------------
        print("\n📐 Calculating column totals:")
        indicator_totals = []

        for header in indicator_headers:
            total = sum_by_header(header)
            indicator_totals.append(total)
            print(f"   ✔ {header} → {total}")

        # --------------------------------------------------
        # Load JSON
        # --------------------------------------------------
        with open(json_path, "r", encoding="utf-8") as f:
            json_data = json.load(f)

        # --------------------------------------------------
        # Update indicators
        # --------------------------------------------------
        print("\n✏️ Updating indicators:")
        for section in json_data:
            if section.get("type") == "data-indicators":
                for i, indicator in enumerate(section.get("indicators", [])):
                    if i >= len(indicator_totals):
                        print(f"⚠️ Skipping {indicator.get('label')} — no corresponding total in Excel")
                        continue
                    old = indicator.get("value")
                    new = indicator_totals[i]
                    indicator["value"] = new
                    print(
                        f"   🔄 {indicator.get('label')} : {old} → {new}"
                  )

        # --------------------------------------------------
        # Write JSON
        # --------------------------------------------------
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(json_data, f, indent=2, ensure_ascii=False)

        print("\n💾 JSON updated successfully")

        # --------------------------------------------------
        # Upload to GCS
        # --------------------------------------------------
        print("\n☁️ Uploading JSON to GCS...")

        gcp_access_path = os.path.join(
            script_dir, "..", "cloud-scripts", "gcp_access.py"
        )
        spec = importlib.util.spec_from_file_location(
            "gcp_access", gcp_access_path
        )
        gcp_access = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(gcp_access)

        folder_url = gcp_access.upload_file_to_gcs_and_get_directory(
            bucket_name=os.environ.get("BUCKET_NAME"),
            source_file_path=json_path,
            destination_blob_name="sg-dashboard/voices-from-the-ground.json",
        )

        if folder_url:
            print(f"Successfully uploaded and got public folder URL: {folder_url}")
        else:
            print("Failed to upload file to GCS. Check logs for details.")

        print("\n🎉 Voices Tab Big Numbers update completed\n")

    except Exception as e:
        print(f"\n❌ Error occurred: {str(e)}\n")
